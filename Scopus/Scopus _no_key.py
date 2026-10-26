# -*- coding: utf-8 -*-
"""通过 Scopus Search API 按日期批量获取题录，并导出 Excel。

这把密钥的实际权限（已核对）：
  - 默认 STANDARD 视图。文档上限：STANDARD 每页 200，COMPLETE 每页 25。
  - 启动时会探测当前密钥实际允许的每页条数；非订户密钥常见上限是 25。
  - 不能用 cursor，单次检索最多取回 5000 条（start 从 0 起）。
  - 周配额见响应头 X-RateLimit-Remaining，用尽后停下，下次接着跑。

不重不漏：
  - 默认按首次入库日 ORIG-LOAD-DATE 逐日检索。每条记录只有一个入库日，日期之间不重叠。
  - 某一天超过 5000 条时，按文献类型、出版年继续拆开，直到每一小段都能取全。
  - 以 EID 为主键去重。每一段取完后核对条数；全部日期结束后再与整段命中数核对。

断点续跑：
  - 进度在输出目录的 scopus_checkpoint.sqlite。
  - 每满 1000 条新记录写一个 Excel；已导出的 EID 记在断点库里，再次运行不会重写。
  - 已完成的日期直接跳过；中断后从上次的页码接着取，已入库的 EID 不再写入。
  - 配额用尽或手动中断后，用同一条命令重新运行即可。

学科筛选写在检索式里，一条记录只要属于其中任一学科就会命中，导出时按 EID 去重。
摘要是可选项，走 Abstract Retrieval API。这把密钥目前不能访问带摘要的 META_ABS 视图。

示例：
  python Scopus.py
  python Scopus.py --year-from 2023 --year-to 2026 --start 2023-01-01 --end 2026-06-30 --out D:\\Scopus题录
  python Scopus.py --subject COMP,MATH --abstract --out D:\\Scopus题录
  python Scopus.py --fill-excel --out D:\\2026年工作\\Scopus题录\\COMP
"""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path

import requests

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("缺少 openpyxl，请先运行: pip install openpyxl")
    raise

# ---------------------------------------------------------------------------
# 基本配置（命令行参数会覆盖这里）
# ---------------------------------------------------------------------------
# 按列表顺序检查。启动时选用检索配额高于保留量的密钥；用尽后自动换下一把。
API_KEYS = [
]
INST_TOKEN = ""

# 检索起止日期（含首尾），格式 YYYY-MM-DD。这里是入库日，不是出版年。
BGDATE = "2023-01-01"
ENDDATE = "2026-06-30"
# 出版年筛选（含首尾）。直接运行时可以改；留在检索式里，避免拆分时从 1800 年扫起。
PUBYEAR_FROM = 2023
PUBYEAR_TO = 2026
# loaddate=首次入库（推荐，不重不漏） | updated=最近装载日 | pubyear=出版年
DATETYPE = "loaddate"
# 在起止日期内怎么切分：day 按日 | month 按月 | year 按年
GRANULARITY = "day"
# 额外检索式，会与日期条件 AND。留空表示该日期内的全部记录。
TERM = ""
# 学科筛选。留空表示不限。可用学科组：理工、生命、医学、社科；也可写代码，如 COMP,ENGI。
SUBJECT = ""
# 是否调用 Abstract Retrieval API 补摘要。默认关闭。
WITH_ABSTRACT = False
OUT_DIR = r"D:\2026年工作\Scopus题录"

SEARCH_URL = "https://api.elsevier.com/content/search/scopus"
ABSTRACT_URL = "https://api.elsevier.com/content/abstract/eid/"
ABSTRACT_VIEW = "META_ABS"
VIEW = "STANDARD"
# 官方文档的每页上限。实际可用条数以启动探测为准。
PAGE_SIZE_BY_VIEW = {
    "STANDARD": 200,
    "COMPLETE": 25,
    "COMPONENT": 25,
}
PAGE_SIZE = PAGE_SIZE_BY_VIEW.get(VIEW.upper(), 25)
OFFSET_LIMIT = 5000
MIN_INTERVAL = 0.15
MAX_RETRIES = 6
REQUEST_TIMEOUT = 120
EXCEL_BATCH_ROWS = 1000
# 剩余周配额低于此值时停下来，避免最后一页把配额打光后丢进度
QUOTA_RESERVE = 1

DATE_FIELDS = {
    "loaddate": "ORIG-LOAD-DATE",
    "updated": "LOAD-DATE",
    "pubyear": "PUBYEAR",
}

# Scopus 文献类型。一条记录只有一个类型，用来把超过 5000 条的结果拆开且不重叠。
DOCTYPES = [
    "ar", "re", "cp", "ch", "bk", "ed", "er", "le", "no", "sh",
    "ab", "cr", "dp", "rp", "pp", "ip",
]

# Scopus 27 个大学科。一条记录可以同时属于多个学科。
SUBJECT_NAMES = {
    "AGRI": "农业与生物科学",
    "ARTS": "艺术与人文",
    "BIOC": "生化、遗传与分子生物学",
    "BUSI": "商业、管理与会计",
    "CENG": "化学工程",
    "CHEM": "化学",
    "COMP": "计算机科学",
    "DECI": "决策科学",
    "DENT": "牙医学",
    "EART": "地球与行星科学",
    "ECON": "经济、计量与金融",
    "ENER": "能源",
    "ENGI": "工程",
    "ENVI": "环境科学",
    "HEAL": "健康专业",
    "IMMU": "免疫与微生物",
    "MATE": "材料科学",
    "MATH": "数学",
    "MEDI": "医学",
    "MULT": "综合性",
    "NEUR": "神经科学",
    "NURS": "护理",
    "PHAR": "药理、毒理与制药",
    "PHYS": "物理与天文",
    "PSYC": "心理学",
    "SOCI": "社会科学",
    "VETE": "兽医学",
}
SUBJECTS = list(SUBJECT_NAMES)
# 与 Scopus 四大门类对应。综合性 MULT 不放入理工，避免把非理工文献一并选入。
SUBJECT_GROUPS = {
    "理工": ["CENG", "CHEM", "COMP", "EART", "ENER", "ENGI", "ENVI", "MATE", "MATH", "PHYS"],
    "生命": ["AGRI", "BIOC", "IMMU", "NEUR", "PHAR"],
    "医学": ["MEDI", "NURS", "HEAL", "DENT", "VETE"],
    "社科": ["ARTS", "BUSI", "DECI", "ECON", "PSYC", "SOCI"],
}

YEAR_MIN = 1800

COLUMNS = [
    "EID",
    "ScopusID",
    "标题",
    "第一作者",
    "作者单位",
    "单位城市",
    "单位国家",
    "期刊",
    "ISSN",
    "EISSN",
    "卷",
    "期",
    "页码",
    "出版日期",
    "出版日期原文",
    "文献类型",
    "来源类型",
    "DOI",
    "PMID",
    "PII",
    "被引次数",
    "开放获取",
    "来源ID",
    "Scopus链接",
    "学科",
    "摘要",
]

OA_LABELS = {
    "0": "非开放获取",
    "1": "金色开放获取",
    "2": "绿色开放获取",
}

ILLEGAL_XLSX = re.compile(r"[\000-\010\013\014\016-\037]")


class QuotaExceeded(RuntimeError):
    """周配额用尽。进度已落库，同一命令可继续。"""


def log(msg: str) -> None:
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)


def parse_date(text: str) -> datetime:
    try:
        return datetime.strptime(text.strip(), "%Y-%m-%d")
    except ValueError as exc:
        raise SystemExit(f"日期格式应为 YYYY-MM-DD，收到: {text}") from exc


def daterange(start: datetime, end: datetime):
    cur = start
    while cur <= end:
        yield cur
        cur += timedelta(days=1)


def month_spans(start: datetime, end: datetime):
    cur = start.replace(day=1)
    while cur <= end:
        if cur.month == 12:
            nxt = cur.replace(year=cur.year + 1, month=1, day=1)
        else:
            nxt = cur.replace(month=cur.month + 1, day=1)
        last = min(end, nxt - timedelta(days=1))
        first = max(start, cur)
        yield first, last
        cur = nxt


def year_spans(start: datetime, end: datetime):
    year = start.year
    while year <= end.year:
        first = max(start, datetime(year, 1, 1))
        last = min(end, datetime(year, 12, 31))
        yield first, last
        year += 1


def day_key(day: datetime) -> str:
    return day.strftime("%Y-%m-%d")


def yyyymmdd(day: datetime) -> str:
    return day.strftime("%Y%m%d")


def clean_text(value) -> str:
    if value is None:
        return ""
    text = " ".join(str(value).replace("\xa0", " ").split())
    return ILLEGAL_XLSX.sub("", text)


def clip_cell(value: str) -> str:
    if len(value) <= 32767:
        return value
    return value[:32760] + "…[截断]"


def join_parts(parts: list[str], sep: str = "; ") -> str:
    return sep.join(p for p in parts if p)


def as_list(value) -> list:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


class RateLimiter:
    def __init__(self, interval: float) -> None:
        self.interval = interval
        self._next = 0.0

    def wait(self) -> None:
        now = time.monotonic()
        if now < self._next:
            time.sleep(self._next - now)
        self._next = time.monotonic() + self.interval


def mask_key(api_key: str) -> str:
    text = api_key.strip()
    if len(text) <= 4:
        return text
    return f"****{text[-4:]}"


class ScopusAPI:
    def __init__(self, api_keys: list[str], inst_token: str) -> None:
        self.api_keys = []
        for key in api_keys:
            text = key.strip()
            if text and text not in self.api_keys:
                self.api_keys.append(text)
        if not self.api_keys:
            raise SystemExit("缺少 Elsevier API Key")
        self.key_index = 0
        self.inst_token = inst_token.strip()
        self.session = requests.Session()
        self.limiter = RateLimiter(MIN_INTERVAL)
        self.quota = {key: {"search": None, "abstract": None} for key in self.api_keys}
        self.view = VIEW.upper()
        self.page_size = PAGE_SIZE_BY_VIEW.get(self.view, PAGE_SIZE)
        self._page_size_ready = False

    @property
    def api_key(self) -> str:
        return self.api_keys[self.key_index]

    def _headers(self) -> dict:
        headers = {
            "X-ELS-APIKey": self.api_key,
            "Accept": "application/json",
        }
        if self.inst_token:
            headers["X-ELS-Insttoken"] = self.inst_token
        return headers

    def _probe_params(self, quota_name: str) -> tuple[str, dict]:
        if quota_name == "abstract":
            return ABSTRACT_URL + "2-s2.0-00000000000", {"view": ABSTRACT_VIEW}
        return SEARCH_URL, {"query": "PUBYEAR IS 1800", "view": self.view, "count": "0", "start": "0"}

    def _probe(self, api_key: str, quota_name: str) -> int:
        """用一次轻量请求读取该密钥的剩余配额。"""
        url, params = self._probe_params(quota_name)
        headers = {"X-ELS-APIKey": api_key, "Accept": "application/json"}
        if self.inst_token:
            headers["X-ELS-Insttoken"] = self.inst_token
        self.limiter.wait()
        try:
            resp = self.session.get(url, headers=headers, params=params, timeout=REQUEST_TIMEOUT)
        except requests.RequestException as exc:
            log(f"密钥 {mask_key(api_key)} 配额查询失败: {exc}")
            self.quota[api_key][quota_name] = 0
            return 0
        raw = resp.headers.get("X-RateLimit-Remaining")
        status = (resp.headers.get("X-ELS-Status") or "").upper()
        try:
            left = int(raw) if raw is not None else 0
        except ValueError:
            left = 0
        if "QUOTA_EXCEEDED" in status:
            left = 0
        self.quota[api_key][quota_name] = left
        log(f"密钥 {mask_key(api_key)} {quota_name} 剩余 {left}")
        return left

    def _ensure_key(self, quota_name: str) -> None:
        current = self.quota[self.api_key][quota_name]
        if current is not None and current > QUOTA_RESERVE:
            return
        best_key = ""
        best_left = QUOTA_RESERVE
        for key in self.api_keys:
            left = self.quota[key][quota_name]
            if left is None:
                left = self._probe(key, quota_name)
            if left > best_left:
                best_key = key
                best_left = left
        if not best_key:
            raise QuotaExceeded(
                f"全部密钥的 {quota_name} 剩余配额都不高于 {QUOTA_RESERVE}。"
                "配额重置后用同一命令继续。"
            )
        if best_key != self.api_key:
            self.key_index = self.api_keys.index(best_key)
            log(f"改用密钥 {mask_key(best_key)}，{quota_name} 剩余 {best_left}")

    def _request(self, url: str, params: dict, label: str, quota_name: str, denied_ok: bool = False) -> dict:
        self._ensure_key(quota_name)
        last_error = ""
        for attempt in range(1, MAX_RETRIES + 1):
            self.limiter.wait()
            try:
                resp = self.session.get(
                    url, headers=self._headers(), params=params, timeout=REQUEST_TIMEOUT
                )
            except requests.RequestException as exc:
                last_error = str(exc)
                self._backoff(attempt, label, last_error)
                continue
            self._read_quota(resp, quota_name)
            if resp.status_code == 200 and resp.content:
                try:
                    return resp.json()
                except json.JSONDecodeError:
                    last_error = "响应不是 JSON"
                    self._backoff(attempt, label, last_error)
                    continue
            status = resp.headers.get("X-ELS-Status") or ""
            last_error = f"HTTP {resp.status_code} {clean_text(status or resp.text[:200])}"
            if resp.status_code == 400 and "translating query" in last_error.lower():
                raise RuntimeError(last_error)
            if resp.status_code in (401, 403) and denied_ok:
                return {"_denied": True, "_status": last_error}
            if quota_name == "abstract" and resp.status_code in (401, 403):
                self.quota[self.api_key][quota_name] = 0
                others = [
                    key for key in self.api_keys
                    if key != self.api_key and self.quota[key][quota_name] != 0
                ]
                if others:
                    log(f"密钥 {mask_key(self.api_key)} 不能访问摘要视图，尝试其他密钥")
                    self._ensure_key(quota_name)
                    continue
                raise SystemExit(
                    "当前密钥不能访问 Abstract Retrieval 的摘要视图 META_ABS。"
                    "学科筛选不依赖该接口，去掉 --abstract 后可以继续下载题录。"
                )
            if "QUOTA_EXCEEDED" in status.upper() or (
                resp.status_code == 429 and "quota" in status.lower()
            ):
                self.quota[self.api_key][quota_name] = 0
                log(f"密钥 {mask_key(self.api_key)} 的 {quota_name} 配额已用尽")
                self._ensure_key(quota_name)
                continue
            if resp.status_code in (429, 500, 502, 503, 504):
                self._backoff(attempt, label, last_error)
                continue
            raise SystemExit(f"{label} 失败: {last_error}")
        raise SystemExit(f"{label} 连续失败 {MAX_RETRIES} 次: {last_error}")

    def _read_quota(self, resp: requests.Response, quota_name: str) -> None:
        raw = resp.headers.get("X-RateLimit-Remaining")
        if raw is None:
            return
        try:
            left = int(raw)
        except ValueError:
            return
        self.quota[self.api_key][quota_name] = left
        if left <= 20:
            log(f"密钥 {mask_key(self.api_key)} {quota_name} 剩余 {left}")

    def _backoff(self, attempt: int, label: str, reason: str) -> None:
        delay = min(60, 2 ** attempt)
        log(f"{label} 第 {attempt} 次失败，{delay}s 后重试: {reason}")
        time.sleep(delay)

    def detect_page_size(self) -> int:
        """按当前视图取文档上限，再向下探测密钥实际允许的每页条数。"""
        if self._page_size_ready:
            return self.page_size
        doc_max = PAGE_SIZE_BY_VIEW.get(self.view, 25)
        if self.view == "STANDARD":
            candidates = [n for n in (doc_max, 100, 50, 25) if n <= doc_max]
        else:
            candidates = [doc_max]
        # 去重并保持从大到小
        seen: list[int] = []
        for n in candidates:
            if n not in seen:
                seen.append(n)
        last_error = ""
        for count in seen:
            self.limiter.wait()
            try:
                resp = self.session.get(
                    SEARCH_URL,
                    headers=self._headers(),
                    params={
                        "query": "PUBYEAR IS 1800",
                        "view": self.view,
                        "count": str(count),
                        "start": "0",
                    },
                    timeout=REQUEST_TIMEOUT,
                )
            except requests.RequestException as exc:
                last_error = str(exc)
                continue
            self._read_quota(resp, "search")
            if resp.status_code == 200:
                self.page_size = count
                self._page_size_ready = True
                log(f"检索视图 {self.view}，每页 {count} 条（文档上限 {doc_max}）")
                return self.page_size
            status = resp.headers.get("X-ELS-Status") or resp.text[:200]
            last_error = f"HTTP {resp.status_code} {clean_text(status)}"
            if resp.status_code == 400 and "maximum number" in last_error.lower():
                continue
            if resp.status_code in (401, 403) and self.view != "STANDARD":
                raise SystemExit(
                    f"当前密钥不能使用视图 {self.view}。改回 VIEW = \"STANDARD\" 后再运行。"
                )
            if resp.status_code == 429 and "QUOTA_EXCEEDED" in last_error:
                self.quota[self.api_key]["search"] = 0
                self._ensure_key("search")
                continue
            raise SystemExit(f"探测每页条数失败: {last_error}")
        raise SystemExit(f"探测每页条数失败: {last_error or '无可用 count'}")

    def search(self, query: str, start: int, count: int) -> dict:
        payload = self._request(
            SEARCH_URL,
            {
                "query": query,
                "view": self.view,
                "start": start,
                "count": count,
                "sort": "+eid",
            },
            f"检索 start={start}",
            "search",
        )
        results = payload.get("search-results")
        if not isinstance(results, dict):
            err = payload.get("service-error") or payload
            raise RuntimeError(f"检索响应异常: {clean_text(json.dumps(err, ensure_ascii=False)[:240])}")
        return results

    def count(self, query: str) -> int:
        results = self.search(query, 0, 0)
        return int(results.get("opensearch:totalResults") or 0)

    def enrich(self, eid: str) -> tuple[str, str]:
        """按官方检索结果里的 dc:description 取摘要，并提取 27 个大学科代码。"""
        try:
            search_payload = self._request(
                SEARCH_URL,
                {"query": f"EID({eid})", "view": "COMPLETE", "count": "1", "start": "0"},
                f"补全 {eid}",
                "search",
                denied_ok=True,
            )
        except QuotaExceeded:
            search_payload = {"_denied": True, "_status": "search quota"}
        subjects, abstract = parse_enrichment(search_payload)
        if abstract and subjects:
            return subjects, abstract
        abstract_payload = self._request(
            ABSTRACT_URL + requests.utils.quote(eid, safe=""),
            {"view": ABSTRACT_VIEW},
            f"补摘要 {eid}",
            "abstract",
            denied_ok=True,
        )
        subjects2, abstract2 = parse_enrichment(abstract_payload)
        subjects = subjects or subjects2
        abstract = abstract or abstract2
        searchable = _has_enrichment_fields(search_payload) or _has_enrichment_fields(abstract_payload)
        if not abstract and not subjects and not searchable:
            raise SystemExit(
                "这些密钥还不能读取摘要。检索 COMPLETE 视图和 Abstract Retrieval 的 META_ABS 都没有返回 dc:description。"
                "Elsevier 开通权限后，用同一命令继续补已保存的 Excel。"
            )
        return subjects, abstract

    def abstract(self, eid: str) -> dict:
        return self._request(
            ABSTRACT_URL + requests.utils.quote(eid, safe=""),
            {"view": ABSTRACT_VIEW},
            f"摘要 {eid}",
            "abstract",
        )


def resolve_subjects(text: str) -> tuple[str, list[str]]:
    """把「理工」或「COMP,ENGI」解析成学科代码。返回 (文件名标记, 代码列表)。"""
    raw = text.strip()
    if not raw:
        return "", []
    codes: list[str] = []
    labels: list[str] = []
    for token in re.split(r"[\s,，;；]+", raw):
        if not token:
            continue
        group = SUBJECT_GROUPS.get(token)
        if group:
            labels.append(token)
            codes.extend(group)
            continue
        code = token.upper()
        if code in SUBJECT_NAMES:
            labels.append(code)
            codes.append(code)
            continue
        known = "、".join(SUBJECT_GROUPS)
        raise SystemExit(f"无法识别学科「{token}」。可用学科组：{known}；或大学科代码如 COMP、ENGI。")
    unique = list(dict.fromkeys(codes))
    return "_".join(labels), unique


def subject_clause(codes: list[str]) -> str:
    if not codes:
        return ""
    return "(" + " OR ".join(f"SUBJAREA({code})" for code in codes) + ")"


def date_span_clause(field: str, start: datetime, end: datetime) -> str:
    """区间只用起止日当天的日期。AFT/BEF 不含当天，不能把边界写成相邻月份。"""
    first = f"{field} IS {yyyymmdd(start)}"
    if start == end:
        return first
    last = f"{field} IS {yyyymmdd(end)}"
    if end - start == timedelta(days=1):
        return f"({first} OR {last})"
    middle = f"{field} AFT {yyyymmdd(start)} AND {field} BEF {yyyymmdd(end)}"
    return f"({first} OR {last} OR ({middle}))"


def pubyear_clause(year_from: int, year_to: int) -> str:
    if year_from == year_to:
        return f"PUBYEAR IS {year_from}"
    return f"PUBYEAR > {year_from - 1} AND PUBYEAR < {year_to + 1}"


def date_query(start: datetime, end: datetime, datetype: str, term: str, subjects: list[str],
               year_from: int, year_to: int) -> str:
    field = DATE_FIELDS[datetype]
    if datetype == "pubyear":
        if start.year == end.year:
            clause = f"PUBYEAR IS {start.year}"
        else:
            clause = f"PUBYEAR > {start.year - 1} AND PUBYEAR < {end.year + 1}"
    else:
        clause = date_span_clause(field, start, end)
    parts = [clause]
    if datetype != "pubyear":
        parts.append(pubyear_clause(year_from, year_to))
    area = subject_clause(subjects)
    if area:
        parts.append(area)
    extra = term.strip()
    if extra:
        parts.append(extra)
    if len(parts) == 1:
        return parts[0]
    return " AND ".join(f"({part})" for part in parts)


def _and(base: str, extra: str) -> str:
    return f"({base}) AND ({extra})"


def plan_shards(api: ScopusAPI, query: str, total: int, start: datetime, end: datetime,
                datetype: str, term: str, subjects: list[str], year_from: int, year_to: int) -> list[tuple[str, int]]:
    """把一次检索拆成若干段，每段不超过 OFFSET_LIMIT，且并集覆盖全部命中。"""
    if total <= OFFSET_LIMIT:
        return [(query, total)]
    if datetype != "pubyear" and start < end:
        log(f"命中 {total}，超过单次 {OFFSET_LIMIT} 条，改按天拆开")
        return _split_days(api, total, start, end, datetype, term, subjects, year_from, year_to)
    log(f"命中 {total}，超过单次 {OFFSET_LIMIT} 条，开始按文献类型拆分")
    shards = _split_doctype(api, query, total)
    flat: list[tuple[str, int]] = []
    for shard_query, shard_total in shards:
        flat.extend(_split_years(api, shard_query, shard_total, year_from, year_to))
    covered = sum(count for _q, count in flat)
    if covered < total:
        raise SystemExit(f"拆分后只有 {covered} 条，原命中 {total}，有遗漏，已停止")
    if covered > total:
        log(f"拆分合计 {covered}，原命中 {total}，多出的部分会在导出时按 EID 去重")
    return [(q, c) for q, c in flat if c > 0]


def _split_days(api: ScopusAPI, total: int, start: datetime, end: datetime,
                datetype: str, term: str, subjects: list[str], year_from: int, year_to: int) -> list[tuple[str, int]]:
    found: list[tuple[str, int]] = []
    covered = 0
    for day in daterange(start, end):
        day_query = date_query(day, day, datetype, term, subjects, year_from, year_to)
        count = api.count(day_query)
        covered += count
        log(f"  {day_key(day)}: {count}")
        if count > OFFSET_LIMIT:
            found.extend(plan_shards(
                api, day_query, count, day, day, datetype, term, subjects, year_from, year_to
            ))
        elif count:
            found.append((day_query, count))
    if covered < total:
        raise SystemExit(f"按天合计 {covered}，原命中 {total}，有遗漏，已停止")
    if covered > total:
        log(f"按天合计 {covered}，原命中 {total}，多出的部分会在导出时按 EID 去重")
    return found


def _split_doctype(api: ScopusAPI, query: str, total: int) -> list[tuple[str, int]]:
    found: list[tuple[str, int]] = []
    used = 0
    for code in DOCTYPES:
        shard = _and(query, f"DOCTYPE({code})")
        try:
            count = api.count(shard)
        except RuntimeError as exc:
            log(f"  文献类型 {code} 不能用于检索，已跳过: {exc}")
            continue
        used += count
        if count:
            log(f"  文献类型 {code}: {count}")
            found.append((shard, count))
    rest = total - used
    if rest > 0:
        excluded = " OR ".join(f"DOCTYPE({code})" for code in DOCTYPES)
        shard = f"{query} AND NOT ({excluded})"
        try:
            count = api.count(shard)
        except RuntimeError as exc:
            log(f"  其他文献类型检索失败，已跳过: {exc}")
            count = 0
        log(f"  其他文献类型: {count}")
        if count:
            found.append((shard, count))
        if count != rest:
            log(f"  类型合计与总量相差 {rest - count}，后续按出版年再拆并按 EID 去重")
    return found or [(query, total)]


def _split_years(api: ScopusAPI, query: str, total: int, lo: int, hi: int) -> list[tuple[str, int]]:
    if total <= OFFSET_LIMIT:
        return [(query, total)]
    if lo >= hi:
        return _split_subject(api, query, total)
    mid = (lo + hi) // 2
    left_q = _and(query, f"PUBYEAR > {lo - 1} AND PUBYEAR < {mid + 1}")
    right_q = _and(query, f"PUBYEAR > {mid} AND PUBYEAR < {hi + 1}")
    left_n = api.count(left_q)
    right_n = api.count(right_q)
    log(f"  出版年 {lo}-{mid}: {left_n}；{mid + 1}-{hi}: {right_n}")
    parts = _split_years(api, left_q, left_n, lo, mid) + _split_years(api, right_q, right_n, mid + 1, hi)
    dated = left_n + right_n
    if dated < total:
        undated_q = _and(query, f"NOT (PUBYEAR > {lo - 1} AND PUBYEAR < {hi + 1})")
        undated_n = api.count(undated_q)
        log(f"  无出版年: {undated_n}")
        if undated_n:
            parts.extend(_split_subject(api, undated_q, undated_n))
    return parts


def _split_subject(api: ScopusAPI, query: str, total: int) -> list[tuple[str, int]]:
    if total <= OFFSET_LIMIT:
        return [(query, total)]
    log(f"  单段仍有 {total} 条，按学科拆分")
    found: list[tuple[str, int]] = []
    for code in SUBJECTS:
        shard = _and(query, f"SUBJAREA({code})")
        count = api.count(shard)
        if not count:
            continue
        if count > OFFSET_LIMIT:
            log(f"  学科 {code} 仍有 {count} 条，改按第一作者首字母拆分")
            if count == total:
                return _split_author(api, query, total)
            found.extend(_split_author(api, shard, count))
            continue
        found.append((shard, count))
    if not found:
        log(f"  学科拆分没有缩小结果，改按第一作者首字母拆分")
        return _split_author(api, query, total)
    return found


def _author_clause(chars: str) -> str:
    return "(" + " OR ".join(f"FIRSTAUTH({ch}*)" for ch in chars) + ")"


def _split_author(api: ScopusAPI, query: str, total: int, chars: str = "ABCDEFGHIJKLMNOPQRSTUVWXYZ") -> list[tuple[str, int]]:
    """按第一作者姓氏首字母对半拆。字母组之间可能重叠，导出时按 EID 去重。"""
    if total <= OFFSET_LIMIT:
        return [(query, total)]
    if len(chars) == 1:
        raise SystemExit(
            f"首字母 {chars} 仍有 {total} 条，超过单次 {OFFSET_LIMIT}。查询: {query}"
        )
    mid = len(chars) // 2
    left, right = chars[:mid], chars[mid:]
    left_q = _and(query, _author_clause(left))
    right_q = _and(query, _author_clause(right))
    left_n = api.count(left_q)
    right_n = api.count(right_q)
    log(f"  作者 {left[0]}-{left[-1]}: {left_n}；{right[0]}-{right[-1]}: {right_n}")
    parts: list[tuple[str, int]] = []
    if left_n:
        parts.extend(_split_author(api, left_q, left_n, left))
    if right_n:
        parts.extend(_split_author(api, right_q, right_n, right))
    if chars == "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
        rest_q = f"{query} AND NOT {_author_clause(chars)}"
        rest_n = api.count(rest_q)
        log(f"  作者首字母非 A-Z: {rest_n}")
        if rest_n > OFFSET_LIMIT:
            raise SystemExit(f"非 A-Z 作者仍有 {rest_n} 条，超过单次 {OFFSET_LIMIT}。查询: {rest_q}")
        if rest_n:
            parts.append((rest_q, rest_n))
    return parts or [(query, total)]


def field_text(entry: dict, key: str) -> str:
    return clean_text(entry.get(key))


def affiliations(entry: dict) -> tuple[str, str, str]:
    names, cities, countries = [], [], []
    for item in as_list(entry.get("affiliation")):
        if not isinstance(item, dict):
            continue
        names.append(clean_text(item.get("affilname")))
        cities.append(clean_text(item.get("affiliation-city")))
        countries.append(clean_text(item.get("affiliation-country")))
    return join_parts(names), join_parts(cities), join_parts(countries)


def open_access(entry: dict) -> str:
    labels = []
    block = entry.get("freetoreadLabel") or {}
    if isinstance(block, dict):
        for item in as_list(block.get("value")):
            if isinstance(item, dict):
                labels.append(clean_text(item.get("$")))
            else:
                labels.append(clean_text(item))
    if labels:
        return join_parts(labels)
    code = field_text(entry, "openaccess")
    if code in OA_LABELS:
        return OA_LABELS[code]
    if entry.get("openaccessFlag") is True:
        return "开放获取"
    if entry.get("openaccessFlag") is False:
        return "非开放获取"
    return code


def scopus_link(entry: dict) -> str:
    for item in as_list(entry.get("link")):
        if isinstance(item, dict) and item.get("@ref") == "scopus":
            return clean_text(item.get("@href"))
    eid = field_text(entry, "eid")
    if eid:
        return f"https://www.scopus.com/record/display.uri?eid={eid}&origin=resultslist"
    return ""


def parse_entry(entry: dict) -> dict | None:
    if not isinstance(entry, dict) or entry.get("error"):
        return None
    eid = field_text(entry, "eid")
    if not eid:
        return None
    identifier = field_text(entry, "dc:identifier")
    scopus_id = identifier.split(":", 1)[-1] if identifier else ""
    affil, city, country = affiliations(entry)
    row = {
        "EID": eid,
        "ScopusID": scopus_id,
        "标题": field_text(entry, "dc:title"),
        "第一作者": field_text(entry, "dc:creator"),
        "作者单位": affil,
        "单位城市": city,
        "单位国家": country,
        "期刊": field_text(entry, "prism:publicationName"),
        "ISSN": field_text(entry, "prism:issn"),
        "EISSN": field_text(entry, "prism:eIssn"),
        "卷": field_text(entry, "prism:volume"),
        "期": field_text(entry, "prism:issueIdentifier"),
        "页码": field_text(entry, "prism:pageRange"),
        "出版日期": field_text(entry, "prism:coverDate"),
        "出版日期原文": field_text(entry, "prism:coverDisplayDate"),
        "文献类型": field_text(entry, "subtypeDescription") or field_text(entry, "subtype"),
        "来源类型": field_text(entry, "prism:aggregationType"),
        "DOI": field_text(entry, "prism:doi"),
        "PMID": field_text(entry, "pubmed-id"),
        "PII": field_text(entry, "pii"),
        "被引次数": field_text(entry, "citedby-count"),
        "开放获取": open_access(entry),
        "来源ID": field_text(entry, "source-id"),
        "Scopus链接": scopus_link(entry),
        "学科": "",
        "摘要": "",
    }
    return {key: clip_cell(row.get(key, "")) for key in COLUMNS}


def parse_page(results: dict) -> list[dict]:
    rows = []
    for entry in as_list(results.get("entry")):
        row = parse_entry(entry)
        if row:
            rows.append(row)
    return rows


class Store:
    def __init__(self, db_path: Path) -> None:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(db_path))
        self.conn.execute("PRAGMA journal_mode=WAL")
        self.conn.execute("PRAGMA synchronous=NORMAL")
        self.conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS jobs (
                job_key TEXT PRIMARY KEY,
                label TEXT NOT NULL,
                query TEXT NOT NULL,
                expected_count INTEGER,
                status TEXT NOT NULL,
                updated_at TEXT
            );
            CREATE TABLE IF NOT EXISTS shards (
                job_key TEXT NOT NULL,
                shard_key TEXT NOT NULL,
                query TEXT NOT NULL,
                expected_count INTEGER NOT NULL,
                next_start INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL,
                PRIMARY KEY (job_key, shard_key)
            );
            CREATE TABLE IF NOT EXISTS records (
                eid TEXT PRIMARY KEY,
                payload TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS abstracts (
                eid TEXT PRIMARY KEY
            );
            CREATE TABLE IF NOT EXISTS job_eids (
                job_key TEXT NOT NULL,
                shard_key TEXT NOT NULL,
                eid TEXT NOT NULL,
                PRIMARY KEY (job_key, shard_key, eid)
            );
            CREATE TABLE IF NOT EXISTS exported (
                eid TEXT PRIMARY KEY,
                job_key TEXT,
                part_no INTEGER
            );
            CREATE TABLE IF NOT EXISTS export_files (
                job_key TEXT NOT NULL,
                part_no INTEGER NOT NULL,
                path TEXT NOT NULL,
                row_count INTEGER NOT NULL,
                PRIMARY KEY (job_key, part_no)
            );
            """
        )
        self._ensure_export_columns()
        self.conn.commit()

    def _ensure_export_columns(self) -> None:
        columns = {row[1] for row in self.conn.execute("PRAGMA table_info(exported)")}
        if "job_key" not in columns:
            self.conn.execute("ALTER TABLE exported ADD COLUMN job_key TEXT")
        if "part_no" not in columns:
            self.conn.execute("ALTER TABLE exported ADD COLUMN part_no INTEGER")

    def close(self) -> None:
        self.conn.close()

    def job(self, job_key: str) -> tuple[str, int | None] | None:
        row = self.conn.execute(
            "SELECT status, expected_count FROM jobs WHERE job_key = ?", (job_key,)
        ).fetchone()
        if not row:
            return None
        return row[0], row[1]

    def mark_job(self, job_key: str, label: str, query: str, status: str, expected: int | None) -> None:
        self.conn.execute(
            """
            INSERT INTO jobs (job_key, label, query, expected_count, status, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(job_key) DO UPDATE SET
                label = excluded.label,
                query = excluded.query,
                expected_count = COALESCE(excluded.expected_count, jobs.expected_count),
                status = excluded.status,
                updated_at = excluded.updated_at
            """,
            (job_key, label, query, expected, status, datetime.now().isoformat(timespec="seconds")),
        )
        self.conn.commit()

    def shards(self, job_key: str) -> list[tuple[str, str, int, int, str]]:
        return self.conn.execute(
            """
            SELECT shard_key, query, expected_count, next_start, status
            FROM shards WHERE job_key = ? ORDER BY shard_key
            """,
            (job_key,),
        ).fetchall()

    def save_shards(self, job_key: str, shards: list[tuple[str, int]]) -> None:
        self.conn.execute("DELETE FROM shards WHERE job_key = ?", (job_key,))
        self.conn.execute("DELETE FROM job_eids WHERE job_key = ?", (job_key,))
        for index, (query, count) in enumerate(shards, start=1):
            self.conn.execute(
                """
                INSERT INTO shards (job_key, shard_key, query, expected_count, next_start, status)
                VALUES (?, ?, ?, ?, 0, 'pending')
                """,
                (job_key, f"{index:03d}", query, count),
            )
        self.conn.commit()

    def mark_shard(self, job_key: str, shard_key: str, status: str, next_start: int | None = None) -> None:
        if next_start is None:
            self.conn.execute(
                "UPDATE shards SET status = ? WHERE job_key = ? AND shard_key = ?",
                (status, job_key, shard_key),
            )
        else:
            self.conn.execute(
                """
                UPDATE shards SET status = ?, next_start = ?
                WHERE job_key = ? AND shard_key = ?
                """,
                (status, next_start, job_key, shard_key),
            )
        self.conn.commit()

    def save_records(self, job_key: str, shard_key: str, rows: list[dict]) -> int:
        new_rows = 0
        for row in rows:
            cur = self.conn.execute(
                "INSERT OR IGNORE INTO records (eid, payload) VALUES (?, ?)",
                (row["EID"], json.dumps(row, ensure_ascii=False)),
            )
            new_rows += cur.rowcount
            self.conn.execute(
                "INSERT OR IGNORE INTO job_eids (job_key, shard_key, eid) VALUES (?, ?, ?)",
                (job_key, shard_key, row["EID"]),
            )
        self.conn.commit()
        return new_rows

    def shard_count(self, job_key: str, shard_key: str) -> int:
        return self.conn.execute(
            "SELECT COUNT(*) FROM job_eids WHERE job_key = ? AND shard_key = ?",
            (job_key, shard_key),
        ).fetchone()[0]

    def distinct_saved(self, job_keys: list[str]) -> int:
        if not job_keys:
            return 0
        marks = ",".join("?" for _ in job_keys)
        return self.conn.execute(
            f"""
            SELECT COUNT(DISTINCT eid) FROM job_eids
            WHERE job_key IN ({marks})
            """,
            job_keys,
        ).fetchone()[0]

    def _export_filter(self, require_abstract: bool) -> str:
        if not require_abstract:
            return ""
        return "AND a.eid IS NOT NULL"

    def unexported_count(self, job_key: str, require_abstract: bool = False) -> int:
        extra = self._export_filter(require_abstract)
        return self.conn.execute(
            f"""
            SELECT COUNT(*)
            FROM (
                SELECT DISTINCT eid FROM job_eids WHERE job_key = ?
            ) j
            JOIN records r ON r.eid = j.eid
            LEFT JOIN exported e ON e.eid = r.eid
            LEFT JOIN abstracts a ON a.eid = r.eid
            WHERE e.eid IS NULL {extra}
            """,
            (job_key,),
        ).fetchone()[0]

    def unexported_rows(self, job_key: str, limit: int, require_abstract: bool = False) -> list[dict]:
        extra = self._export_filter(require_abstract)
        rows = self.conn.execute(
            f"""
            SELECT r.eid, r.payload
            FROM (
                SELECT DISTINCT eid FROM job_eids WHERE job_key = ?
            ) j
            JOIN records r ON r.eid = j.eid
            LEFT JOIN exported e ON e.eid = r.eid
            LEFT JOIN abstracts a ON a.eid = r.eid
            WHERE e.eid IS NULL {extra}
            ORDER BY r.eid
            LIMIT ?
            """,
            (job_key, limit),
        ).fetchall()
        return [json.loads(payload) for _eid, payload in rows]

    def pending_abstract_eids(self, job_key: str, limit: int) -> list[str]:
        rows = self.conn.execute(
            """
            SELECT j.eid
            FROM (
                SELECT DISTINCT eid FROM job_eids WHERE job_key = ?
            ) j
            LEFT JOIN abstracts a ON a.eid = j.eid
            WHERE a.eid IS NULL
            ORDER BY j.eid
            LIMIT ?
            """,
            (job_key, limit),
        ).fetchall()
        return [row[0] for row in rows]

    def save_abstract(self, eid: str, subject: str, abstract: str) -> None:
        row = self.conn.execute("SELECT payload FROM records WHERE eid = ?", (eid,)).fetchone()
        if row is None:
            return
        payload = json.loads(row[0])
        payload["学科"] = clip_cell(subject)
        payload["摘要"] = clip_cell(abstract)
        self.conn.execute(
            "UPDATE records SET payload = ? WHERE eid = ?",
            (json.dumps(payload, ensure_ascii=False), eid),
        )
        self.conn.execute("INSERT OR IGNORE INTO abstracts (eid) VALUES (?)", (eid,))
        self.conn.commit()

    def next_part_no(self, job_key: str) -> int:
        row = self.conn.execute(
            "SELECT MAX(part_no) FROM export_files WHERE job_key = ?",
            (job_key,),
        ).fetchone()
        return (row[0] or 0) + 1

    def mark_exported(self, job_key: str, part_no: int, path: Path, eids: list[str]) -> None:
        self.conn.execute(
            """
            INSERT INTO export_files (job_key, part_no, path, row_count)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(job_key, part_no) DO UPDATE SET
                path = excluded.path,
                row_count = excluded.row_count
            """,
            (job_key, part_no, str(path), len(eids)),
        )
        self.conn.executemany(
            "INSERT OR IGNORE INTO exported (eid, job_key, part_no) VALUES (?, ?, ?)",
            ((eid, job_key, part_no) for eid in eids),
        )
        self.conn.commit()


ASJC_MAJOR = {
    10: "MULT", 11: "AGRI", 12: "ARTS", 13: "BIOC", 14: "BUSI", 15: "CENG",
    16: "CHEM", 17: "COMP", 18: "DECI", 19: "EART", 20: "ECON", 21: "ENER",
    22: "ENGI", 23: "ENVI", 24: "IMMU", 25: "MATE", 26: "MATH", 27: "MEDI",
    28: "NEUR", 29: "NURS", 30: "PHAR", 31: "PHYS", 32: "PSYC", 33: "SOCI",
    34: "VETE", 35: "DENT", 36: "HEAL",
}


def major_subject_codes(payload: dict) -> str:
    """从检索条目或摘要响应中取出 Scopus 27 个大学科代码，如 COMP。"""
    found: list[str] = []

    def add(code: str) -> None:
        if code in SUBJECT_NAMES and code not in found:
            found.append(code)

    def walk(node) -> None:
        if isinstance(node, dict):
            abbrev = clean_text(node.get("@abbrev") or node.get("abbrev")).upper()
            add(abbrev)
            raw = clean_text(node.get("@code") or node.get("code"))
            if raw.isdigit():
                add(ASJC_MAJOR.get(int(raw) // 100, ""))
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(payload)
    return "; ".join(found)


def description_text(payload: dict) -> str:
    """官方示例把摘要放在 dc:description。"""
    found: list[str] = []

    def walk(node) -> None:
        if isinstance(node, dict):
            if "dc:description" in node:
                text = abstract_text(node.get("dc:description"))
                if text and text not in found:
                    found.append(text)
            for key, value in node.items():
                if key != "dc:description":
                    walk(value)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    walk(payload)
    return found[0] if found else ""


def _has_enrichment_fields(payload: dict) -> bool:
    if not isinstance(payload, dict) or payload.get("_denied"):
        return False
    text = json.dumps(payload, ensure_ascii=False)
    return "dc:description" in text or "subject-area" in text


def parse_enrichment(payload: dict) -> tuple[str, str]:
    if not isinstance(payload, dict) or payload.get("_denied"):
        return "", ""
    return major_subject_codes(payload), description_text(payload)


def abstract_text(node) -> str:
    if node is None:
        return ""
    if isinstance(node, str):
        return clean_text(node)
    if isinstance(node, dict):
        return clean_text(node.get("$") or node.get("#text") or "")
    return clean_text(str(node))


def parse_abstract(payload: dict) -> tuple[str, str]:
    body = payload.get("abstracts-retrieval-response") or {}
    core = body.get("coredata") or {}
    description = abstract_text(core.get("dc:description"))
    block = body.get("subject-areas") or {}
    areas = []
    items = block.get("subject-area") if isinstance(block, dict) else block
    for item in as_list(items):
        if isinstance(item, dict):
            name = abstract_text(item.get("$")) or clean_text(item.get("@abbrev"))
            abbrev = clean_text(item.get("@abbrev"))
            if name and abbrev and abbrev not in name:
                areas.append(f"{name}({abbrev})")
            elif name:
                areas.append(name)
        else:
            areas.append(clean_text(item))
    return join_parts(areas, " | "), description


def fill_abstracts(api: ScopusAPI, store: Store, job_key: str) -> None:
    while True:
        eids = store.pending_abstract_eids(job_key, PAGE_SIZE)
        if not eids:
            return
        for eid in eids:
            payload = api.abstract(eid)
            subject, abstract = parse_abstract(payload)
            store.save_abstract(eid, subject, abstract)
        log(f"  摘要已补 {len(eids)} 条")


def fetch_shard(api: ScopusAPI, store: Store, out_dir: Path, label: str, with_abstract: bool,
                job_key: str, shard_key: str, query: str, expected: int, next_start: int) -> None:
    if expected == 0:
        store.mark_shard(job_key, shard_key, "done", 0)
        return
    start = next_start
    log(f"  分段 {shard_key} 共 {expected} 条，从第 {start + 1} 条继续")
    while start < expected:
        if start >= OFFSET_LIMIT:
            raise SystemExit(f"分段 {shard_key} 超过偏移上限 {OFFSET_LIMIT}: {query}")
        page_size = min(api.page_size, expected - start)
        results = api.search(query, start, page_size)
        rows = parse_page(results)
        if len(rows) < page_size:
            raise RuntimeError(
                f"分段 {shard_key} 第 {start} 条起期望 {page_size} 条，实际 {len(rows)}"
            )
        store.save_records(job_key, shard_key, rows)
        start += page_size
        store.mark_shard(job_key, shard_key, "paging", start)
        log(f"  分段 {shard_key} 进度 {start}/{expected}")
        if with_abstract:
            fill_abstracts(api, store, job_key)
        flush_excel(store, job_key, out_dir, label, final=False, require_abstract=with_abstract)
    saved = store.shard_count(job_key, shard_key)
    if saved != expected:
        store.mark_shard(job_key, shard_key, "pending", 0)
        raise RuntimeError(
            f"分段 {shard_key} 核对失败: 应有 {expected}，已保存 {saved}，将按 EID 从头再取缺失记录"
        )
    store.mark_shard(job_key, shard_key, "done", start)


def excel_path(out_dir: Path, label: str, part_no: int) -> Path:
    safe = re.sub(r'[<>:"/\\|?*]+', "_", label).strip(" ._")
    return out_dir / f"Scopus_{safe}_{part_no}.xlsx"


def write_excel_file(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".xlsx.part")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("题录")
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])
    wb.save(tmp)
    tmp.replace(path)


def flush_excel(store: Store, job_key: str, out_dir: Path, label: str, final: bool,
                require_abstract: bool = False) -> int:
    """每满 1000 条写一个文件。任务结束时把不足 1000 的余数也写完。

    按 EID 排序取下一批，文件序号记在断点库。中断后同一批会写回同一个文件，不会另存一份。
    """
    written = 0
    while True:
        pending = store.unexported_count(job_key, require_abstract)
        if pending == 0:
            return written
        if not final and pending < EXCEL_BATCH_ROWS:
            return written
        take = EXCEL_BATCH_ROWS if pending >= EXCEL_BATCH_ROWS else pending
        rows = store.unexported_rows(job_key, take, require_abstract)
        if not rows:
            return written
        part_no = store.next_part_no(job_key)
        path = excel_path(out_dir, label, part_no)
        write_excel_file(path, rows)
        store.mark_exported(job_key, part_no, path, [row["EID"] for row in rows])
        written += len(rows)
        log(f"已写入 {path}（{len(rows)} 行）")
        if len(rows) < EXCEL_BATCH_ROWS:
            return written


def run_unit(api: ScopusAPI, store: Store, out_dir: Path, label: str, query: str,
            span_start: datetime, span_end: datetime, datetype: str, term: str,
            subjects: list[str], year_from: int, year_to: int, with_abstract: bool) -> str:
    job_key = label
    current = store.job(job_key)
    if current and current[0] == "done":
        log(f"已完成，跳过: {label}")
        return job_key
    if current and current[0] == "ready":
        shards = store.shards(job_key)
        expected = current[1] or 0
    else:
        log(f"检索: {label}")
        log(f"  {query}")
        expected = api.count(query)
        log(f"  命中 {expected}")
        store.mark_job(job_key, label, query, "planning", expected)
        planned = plan_shards(
            api, query, expected, span_start, span_end, datetype, term, subjects, year_from, year_to
        ) if expected else []
        store.save_shards(job_key, planned)
        store.mark_job(job_key, label, query, "ready", expected)
        shards = store.shards(job_key)
    for shard_key, shard_query, shard_total, next_start, status in shards:
        if status == "done":
            log(f"  分段 {shard_key} 已完成，跳过（{shard_total} 条）")
            continue
        last_error = ""
        for attempt in range(1, MAX_RETRIES + 1):
            try:
                fetch_shard(
                    api, store, out_dir, label, with_abstract, job_key, shard_key,
                    shard_query, shard_total, next_start,
                )
                break
            except QuotaExceeded:
                raise
            except RuntimeError as exc:
                last_error = str(exc)
                log(f"  {last_error}，重取本页（{attempt}/{MAX_RETRIES}）")
                time.sleep(min(30, 2 ** attempt))
                row = store.conn.execute(
                    "SELECT next_start FROM shards WHERE job_key = ? AND shard_key = ?",
                    (job_key, shard_key),
                ).fetchone()
                next_start = row[0] if row else next_start
        else:
            raise SystemExit(f"{label} 分段 {shard_key} 未取全: {last_error}")
    saved = store.distinct_saved([job_key])
    if saved != expected:
        raise SystemExit(f"{label} 核对失败: 应有 {expected}，已保存唯一 EID {saved}")
    if with_abstract:
        fill_abstracts(api, store, job_key)
    written = flush_excel(store, job_key, out_dir, label, final=True, require_abstract=with_abstract)
    store.mark_job(job_key, label, query, "done", expected)
    log(f"完成: {label}，EID {expected}，新写入 Excel {written} 行")
    return job_key


class EnrichStore:
    def __init__(self, db_path: Path) -> None:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self.conn = sqlite3.connect(str(db_path))
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS enrich (
                eid TEXT PRIMARY KEY,
                subjects TEXT NOT NULL,
                abstract TEXT NOT NULL
            )
            """
        )
        self.conn.commit()

    def get(self, eid: str) -> tuple[str, str] | None:
        row = self.conn.execute(
            "SELECT subjects, abstract FROM enrich WHERE eid = ?", (eid,)
        ).fetchone()
        return (row[0], row[1]) if row else None

    def save(self, eid: str, subjects: str, abstract: str) -> None:
        self.conn.execute(
            "INSERT OR REPLACE INTO enrich (eid, subjects, abstract) VALUES (?, ?, ?)",
            (eid, subjects, abstract),
        )
        self.conn.commit()

    def close(self) -> None:
        self.conn.close()


def fill_saved_excels(api: ScopusAPI, out_dir: Path) -> None:
    """为已经导出的 Excel 补「学科」和「摘要」。同一 EID 只请求一次。"""
    files = sorted(
        path for path in out_dir.rglob("*.xlsx")
        if not path.name.startswith("~$") and ".part" not in path.name
    )
    if not files:
        raise SystemExit(f"目录里没有 Excel: {out_dir}")
    store = EnrichStore(out_dir / "excel_enrich.sqlite")
    filled = 0
    try:
        for path in files:
            filled += _fill_one_excel(api, store, path)
    except QuotaExceeded as exc:
        log(str(exc))
    finally:
        store.close()
    log(f"摘要和学科补写完成，本次新请求 {filled} 条。断点: {out_dir / 'excel_enrich.sqlite'}")


def _fill_one_excel(api: ScopusAPI, store: EnrichStore, path: Path) -> int:
    try:
        wb = load_workbook(path)
    except PermissionError:
        log(f"文件被占用，跳过: {path}")
        return 0
    ws = wb.active
    header = [clean_text(cell.value) for cell in ws[1]]
    if "EID" not in header:
        wb.close()
        return 0
    for name in ("学科", "摘要"):
        if name not in header:
            ws.cell(1, len(header) + 1, name)
            header.append(name)
    eid_col = header.index("EID") + 1
    subject_col = header.index("学科") + 1
    abstract_col = header.index("摘要") + 1
    filled = 0
    changed = False
    for row in range(2, ws.max_row + 1):
        eid = clean_text(ws.cell(row, eid_col).value)
        if not eid:
            continue
        has_subject = bool(clean_text(ws.cell(row, subject_col).value))
        has_abstract = bool(clean_text(ws.cell(row, abstract_col).value))
        if has_subject and has_abstract:
            continue
        cached = store.get(eid)
        if cached is None:
            subjects, abstract = api.enrich(eid)
            store.save(eid, subjects, abstract)
            filled += 1
        else:
            subjects, abstract = cached
        if not has_subject and subjects:
            ws.cell(row, subject_col, clip_cell(subjects))
            changed = True
        if not has_abstract and abstract:
            ws.cell(row, abstract_col, clip_cell(abstract))
            changed = True
        if filled and filled % 20 == 0 and changed:
            _save_workbook(wb, path)
            changed = False
            log(f"  {path.name} 已补 {filled} 条")
    if changed:
        _save_workbook(wb, path)
    wb.close()
    log(f"{path.name} 补写 {filled} 条")
    return filled


def _save_workbook(wb, path: Path) -> None:
    tmp = path.with_suffix(".xlsx.part")
    wb.save(tmp)
    tmp.replace(path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="按日期从 Scopus 下载题录并导出 Excel")
    parser.add_argument("--start", default=BGDATE, help="开始日期 YYYY-MM-DD，含当天")
    parser.add_argument("--end", default=ENDDATE, help="结束日期 YYYY-MM-DD，含当天")
    parser.add_argument("--datetype", default=DATETYPE, choices=sorted(DATE_FIELDS),
                        help="loaddate 首次入库（默认，不重不漏）；updated 最近装载日；pubyear 出版年")
    parser.add_argument("--granularity", default=GRANULARITY, choices=("day", "month", "year"),
                        help="在起止日期内按日、按月或按年切分。直接运行脚本时会改为交互选择")
    parser.add_argument("--term", default=TERM, help="附加检索式，与日期条件 AND")
    parser.add_argument("--year-from", type=int, default=PUBYEAR_FROM, help="出版年起，含该年")
    parser.add_argument("--year-to", type=int, default=PUBYEAR_TO, help="出版年止，含该年")
    parser.add_argument("--subject", default=SUBJECT,
                        help="学科筛选。学科组：理工、生命、医学、社科；或代码如 COMP,ENGI。留空不限")
    parser.add_argument("--abstract", action="store_true", default=WITH_ABSTRACT,
                        help="调用 Abstract Retrieval API 补摘要和学科名称")
    parser.add_argument("--fill-excel", action="store_true",
                        help="只为已保存的 Excel 补摘要和学科代码，不重新下载")
    parser.add_argument("--out", default=OUT_DIR, help="Excel 与断点库的保存目录")
    parser.add_argument("--api-key", default="", help="密钥，多个用逗号分隔。留空则使用脚本里的 API_KEYS")
    parser.add_argument("--inst-token", default=INST_TOKEN, help="机构令牌，没有就留空")
    return parser


def unit_label(base: str, subject_label: str, year_from: int, year_to: int, with_abstract: bool) -> str:
    parts = [base, f"y{year_from}-{year_to}"]
    if subject_label:
        parts.append(subject_label)
    if with_abstract:
        parts.append("abs")
    return "_".join(parts)


def is_full_month(first: datetime, last: datetime) -> bool:
    if first.day != 1 or first.year != last.year or first.month != last.month:
        return False
    if first.month == 12:
        month_end = datetime(first.year, 12, 31)
    else:
        month_end = datetime(first.year, first.month + 1, 1) - timedelta(days=1)
    return last == month_end


def span_base(first: datetime, last: datetime, datetype: str, granularity: str) -> str:
    if datetype == "pubyear":
        return f"pubyear_{first.year}"
    if granularity == "day" and first == last:
        return f"{day_key(first)}_{datetype}"
    if granularity == "month" and is_full_month(first, last):
        return f"{first:%Y-%m}_month_{datetype}"
    if granularity == "year" and first == datetime(first.year, 1, 1) and last == datetime(first.year, 12, 31):
        return f"{first.year}_year_{datetype}"
    return f"{day_key(first)}_{day_key(last)}_{granularity}_{datetype}"


def units(start: datetime, end: datetime, datetype: str, granularity: str, term: str,
          subjects: list[str], subject_label: str, year_from: int, year_to: int,
          with_abstract: bool) -> list[tuple]:
    if datetype == "pubyear":
        spans = ((datetime(year, 1, 1), datetime(year, 12, 31)) for year in range(start.year, end.year + 1))
        use = "year"
    elif granularity == "month":
        spans = month_spans(start, end)
        use = granularity
    elif granularity == "year":
        spans = year_spans(start, end)
        use = granularity
    else:
        spans = ((day, day) for day in daterange(start, end))
        use = "day"
    items = []
    for first, last in spans:
        label = unit_label(span_base(first, last, datetype, use), subject_label, year_from, year_to, with_abstract)
        items.append((label, date_query(first, last, datetype, term, subjects, year_from, year_to), first, last))
    return items


def ask_line(prompt: str) -> str:
    try:
        return input(prompt).strip().strip('"')
    except EOFError as exc:
        raise SystemExit("未读到输入，已退出。") from exc


def prompt_settings(args: argparse.Namespace) -> None:
    """不带参数直接运行时，依次选择日期范围、粒度、学科和导出路径。"""
    print("Scopus 题录下载")
    print("直接回车表示使用括号里的默认值。")
    while True:
        start_text = ask_line(f"开始日期 YYYY-MM-DD [{args.start}]: ") or args.start
        end_text = ask_line(f"结束日期 YYYY-MM-DD [{args.end}]: ") or args.end
        try:
            start = parse_date(start_text)
            end = parse_date(end_text)
        except SystemExit as exc:
            print(exc)
            continue
        if end < start:
            print("结束日期不能早于开始日期。")
            continue
        args.start = day_key(start)
        args.end = day_key(end)
        break
    grain_names = {"1": "day", "2": "month", "3": "year", "day": "day", "month": "month", "year": "year", "日": "day", "月": "month", "年": "year"}
    while True:
        choice = ask_line("检索粒度 1按日  2按月  3按年 [1]: ") or "1"
        grain = grain_names.get(choice)
        if grain:
            args.granularity = grain
            break
        print("请输入 1、2 或 3。")
    while True:
        default_years = f"{args.year_from}-{args.year_to}"
        text = ask_line(f"出版年 起-止 [{default_years}]: ") or default_years
        matched = re.fullmatch(r"(\d{4})\s*[-~到至]\s*(\d{4})", text)
        if not matched:
            print("请按 2023-2026 这种格式输入。")
            continue
        year_from, year_to = int(matched.group(1)), int(matched.group(2))
        if year_to < year_from:
            print("结束出版年不能早于开始出版年。")
            continue
        args.year_from = year_from
        args.year_to = year_to
        break
    print("学科：0 不限  1 理工  2 生命  3 医学  4 社科")
    print("也可直接输入学科代码，多个用逗号分隔，例如 COMP,ENGI")
    subject_menu = {"0": "", "1": "理工", "2": "生命", "3": "医学", "4": "社科"}
    while True:
        choice = ask_line("学科 [0]: ") or "0"
        subject = subject_menu.get(choice, choice)
        try:
            resolve_subjects(subject)
        except SystemExit as exc:
            print(exc)
            continue
        args.subject = subject
        break
    path = ask_line(f"导出路径，可留空 [{args.out}]: ")
    if path:
        args.out = path


def parse_api_keys(text: str) -> list[str]:
    if not text.strip():
        return list(API_KEYS)
    keys = [part.strip() for part in re.split(r"[\s,，;；]+", text) if part.strip()]
    if not keys:
        raise SystemExit("缺少 Elsevier API Key")
    return keys


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    args = build_parser().parse_args()
    if len(sys.argv) == 1:
        prompt_settings(args)
    start = parse_date(args.start)
    end = parse_date(args.end)
    if end < start:
        raise SystemExit("结束日期不能早于开始日期")
    if args.year_to < args.year_from:
        raise SystemExit("结束出版年不能早于开始出版年")
    api_keys = parse_api_keys(args.api_key)
    subject_label, subjects = resolve_subjects(args.subject)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    db_path = out_dir / "scopus_checkpoint.sqlite"
    grain_text = {"day": "按日", "month": "按月", "year": "按年"}[args.granularity]
    log(f"日期 {day_key(start)} ~ {day_key(end)}，类型 {args.datetype}，{grain_text}")
    log(f"出版年: {args.year_from}-{args.year_to}")
    if subjects:
        names = "、".join(f"{code} {SUBJECT_NAMES[code]}" for code in subjects)
        log(f"学科筛选: {names}")
    else:
        log("学科筛选: 不限")
    log(f"摘要接口: {'开启' if args.abstract else '关闭'}")
    log(f"输出目录: {out_dir}")
    if args.datetype != "loaddate":
        log("当前不是首次入库日。出版年或最近装载日可能与入库日口径不同，脚本仍按 EID 去重并核对条数。")
    api = ScopusAPI(api_keys, args.inst_token)
    if args.fill_excel:
        fill_saved_excels(api, out_dir)
        return
    api._ensure_key("search")
    log(f"使用密钥 {mask_key(api.api_key)}，检索剩余 {api.quota[api.api_key]['search']}")
    api.detect_page_size()
    store = Store(db_path)
    try:
        job_keys = []
        for label, query, span_start, span_end in units(
            start, end, args.datetype, args.granularity, args.term, subjects, subject_label,
            args.year_from, args.year_to, args.abstract,
        ):
            job_keys.append(run_unit(
                api, store, out_dir, label, query, span_start, span_end,
                args.datetype, args.term, subjects, args.year_from, args.year_to, args.abstract,
            ))
        whole_query = date_query(start, end, args.datetype, args.term, subjects, args.year_from, args.year_to)
        whole = api.count(whole_query)
        saved = store.distinct_saved(job_keys)
        log(f"核对: 整段命中 {whole}，已保存唯一 EID {saved}")
        if saved < whole:
            raise SystemExit(f"结果不完整: 命中 {whole}，已保存 {saved}。重新运行将只补缺失部分。")
        if saved > whole:
            log(f"唯一 EID 比整段命中多 {saved - whole}。入库日逐日不应出现这种情况，请检查是否换过检索式却沿用了旧断点库。")
        exported = store.conn.execute("SELECT COUNT(*) FROM exported").fetchone()[0]
        search_left = api.quota[api.api_key]["search"]
        abstract_left = api.quota[api.api_key]["abstract"]
        if search_left is None:
            search_left = "未知"
        if abstract_left is None:
            abstract_left = "未使用"
        log(f"全部完成。已导出不重复 EID {exported} 条。检索剩余配额 {search_left}，摘要剩余配额 {abstract_left}。断点库: {db_path}")
    except QuotaExceeded as exc:
        log(str(exc))
        raise SystemExit(0) from exc
    finally:
        store.close()


if __name__ == "__main__":
    main()
