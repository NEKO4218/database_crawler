# -*- coding: utf-8 -*-
"""从题录 Excel 中读取 PMCID，下载 PMC 开放获取全文（PDF / XML）。

只处理有 PMCID 的行。按「单个 Excel」流式读取并边读边下，不把全部表一次读进内存。
开放获取全文优先走 EuropePMC：
  - PDF: https://europepmc.org/api/getPdf?pmcid=PMC...
  - XML: https://www.ebi.ac.uk/europepmc/webservices/rest/PMC.../fullTextXML
XML 若失败，再回退 NCBI EFetch（db=pmc）。

不请求出版社付费全文；没有开放获取文件时记为 no_oa，不中断整批。

示例：
  python download_pmc_fulltext.py --excel D:\\PubMed题录\\PubMed_2023-01-02_2023-01-02_pdat.xlsx --out D:\\PubMed全文 --format pdf --resume
  python download_pmc_fulltext.py --excel-dir D:\\PubMed题录 --out D:\\PubMed全文 --format pdf --resume
  python download_pmc_fulltext.py --excel-dir D:\\PubMed题录 --format both --limit 100
"""

from __future__ import annotations

import argparse
import csv
import re
import sqlite3
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Iterator

import requests

try:
    from openpyxl import load_workbook
except ImportError:
    print("缺少 openpyxl，请先运行: pip install openpyxl")
    raise

EUROPEPMC_PDF = "https://europepmc.org/api/getPdf"
EUROPEPMC_XML = "https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/fullTextXML"
EFETCH_URL = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi"

DEFAULT_EXCEL_DIR = r"D:\2026年工作\PubMed题录"
DEFAULT_OUT_DIR = ""
API_KEY = "4fef42571ead577458ba2a15d8ea31eeaf09"
EMAIL = ""
TOOL = "download_pmc_fulltext.py"

MIN_INTERVAL = 0.35
REQUEST_TIMEOUT = 120
MAX_RETRIES = 5
# 流式读取时，每攒够一批再下载，避免单文件亿级中间集合
BATCH_SIZE = 200
EXCEL_SUFFIXES = {".xlsx", ".xls"}
PMCID_COL_CANDIDATES = (
    "PMCID",
    "pmcid",
    "Pmcid",
    "PMC ID",
    "PMC",
)

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (compatible; download_pmc_fulltext/1.0; "
        "mailto:research@example.edu)"
    ),
    "Accept": "*/*",
}


def log(msg: str) -> None:
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)


def normalize_pmcid(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null", "-"}:
        return ""
    match = re.search(r"(?:PMC)?\s*(\d+)", text, flags=re.I)
    if not match:
        return ""
    return f"PMC{match.group(1)}"


def is_pdf(data: bytes) -> bool:
    return bool(data) and data[:5] == b"%PDF-"


def is_xml(data: bytes) -> bool:
    if not data:
        return False
    head = data.lstrip()[:200].lower()
    return head.startswith(b"<?xml") or head.startswith(b"<")


class RateLimiter:
    def __init__(self, interval: float) -> None:
        self.interval = interval
        self._next = 0.0

    def wait(self) -> None:
        now = time.monotonic()
        if now < self._next:
            time.sleep(self._next - now)
        self._next = time.monotonic() + self.interval


class Store:
    def __init__(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.path = path
        self.conn = sqlite3.connect(str(path))
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS downloads (
                pmcid TEXT NOT NULL,
                kind TEXT NOT NULL,
                status TEXT NOT NULL,
                path TEXT,
                detail TEXT,
                updated_at TEXT,
                PRIMARY KEY (pmcid, kind)
            )
            """
        )
        self.conn.commit()

    def close(self) -> None:
        self.conn.close()

    def status(self, pmcid: str, kind: str) -> str | None:
        row = self.conn.execute(
            "SELECT status FROM downloads WHERE pmcid = ? AND kind = ?",
            (pmcid, kind),
        ).fetchone()
        return row[0] if row else None

    def mark(self, pmcid: str, kind: str, status: str, path: str = "", detail: str = "") -> None:
        self.conn.execute(
            """
            INSERT INTO downloads (pmcid, kind, status, path, detail, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(pmcid, kind) DO UPDATE SET
                status = excluded.status,
                path = excluded.path,
                detail = excluded.detail,
                updated_at = excluded.updated_at
            """,
            (
                pmcid,
                kind,
                status,
                path,
                detail[:500],
                datetime.now().isoformat(timespec="seconds"),
            ),
        )
        self.conn.commit()


def list_excels(folder: Path) -> list[Path]:
    files = []
    for path in sorted(folder.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() not in EXCEL_SUFFIXES:
            continue
        if path.name.startswith("~$") or path.name.startswith("."):
            continue
        files.append(path)
    return files


def resolve_excel_inputs(excel: str, excel_dir: str) -> list[Path]:
    """优先 --excel 单文件；否则扫描 --excel-dir。"""
    excel_text = (excel or "").strip().strip('"')
    if excel_text:
        path = Path(excel_text)
        if not path.is_file():
            raise SystemExit(f"Excel 文件不存在: {path}")
        if path.suffix.lower() not in EXCEL_SUFFIXES:
            raise SystemExit(f"不是 Excel 文件: {path}")
        return [path]
    folder = Path((excel_dir or "").strip().strip('"') or DEFAULT_EXCEL_DIR)
    if not folder.is_dir():
        raise SystemExit(f"题录目录不存在: {folder}")
    files = list_excels(folder)
    if not files:
        raise SystemExit(f"未找到 Excel：{folder}")
    return files


def find_pmcid_column(columns) -> str | None:
    names = [str(c).strip() if c is not None else "" for c in columns]
    lower_map = {name.lower(): name for name in names if name}
    for cand in PMCID_COL_CANDIDATES:
        key = cand.lower().strip()
        if key in lower_map:
            return lower_map[key]
    for name in names:
        if name and "pmcid" in name.lower().replace(" ", ""):
            return name
    return None


def iter_pmcids_xlsx(path: Path) -> Iterator[str]:
    """只读流式扫描 .xlsx，按行产出 PMCID，本文件内去重。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return
        col_name = find_pmcid_column(header or ())
        if not col_name:
            log(f"跳过无 PMCID 列的表: {path.name}")
            return
        headers = [str(c).strip() if c is not None else "" for c in (header or ())]
        try:
            idx = headers.index(col_name)
        except ValueError:
            log(f"跳过无法定位 PMCID 列的表: {path.name}")
            return
        seen: set[str] = set()
        for row in rows:
            if not row or idx >= len(row):
                continue
            pmcid = normalize_pmcid(row[idx])
            if not pmcid or pmcid in seen:
                continue
            seen.add(pmcid)
            yield pmcid
    finally:
        wb.close()


def iter_pmcids_xls(path: Path) -> Iterator[str]:
    """旧版 .xls：用 pandas 分块读（若引擎支持）；否则整表一次。"""
    try:
        import pandas as pd
    except ImportError as exc:
        raise SystemExit("读取 .xls 需要 pandas / xlrd，请先安装") from exc
    seen: set[str] = set()
    try:
        chunks = pd.read_excel(path, dtype=str, chunksize=5000)
    except TypeError:
        chunks = [pd.read_excel(path, dtype=str)]
    first = True
    col = None
    for chunk in chunks:
        if first:
            col = find_pmcid_column(chunk.columns)
            first = False
            if not col:
                log(f"跳过无 PMCID 列的表: {path.name}")
                return
        for value in chunk[col].tolist():
            pmcid = normalize_pmcid(value)
            if not pmcid or pmcid in seen:
                continue
            seen.add(pmcid)
            yield pmcid


def iter_pmcids(path: Path) -> Iterator[str]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        yield from iter_pmcids_xlsx(path)
    elif suffix == ".xls":
        yield from iter_pmcids_xls(path)
    else:
        log(f"跳过不支持的后缀: {path.name}")


def request_bytes(session: requests.Session, limiter: RateLimiter, url: str,
                  params: dict | None = None) -> tuple[int, bytes, str]:
    last = ""
    for attempt in range(1, MAX_RETRIES + 1):
        limiter.wait()
        try:
            resp = session.get(
                url,
                params=params,
                timeout=REQUEST_TIMEOUT,
                allow_redirects=True,
                headers=HEADERS,
            )
        except requests.RequestException as exc:
            last = str(exc)
            time.sleep(min(60, 2 ** attempt))
            continue
        ctype = (resp.headers.get("Content-Type") or "").lower()
        if resp.status_code == 200 and resp.content:
            return resp.status_code, resp.content, ctype
        if resp.status_code in (404, 410):
            return resp.status_code, resp.content or b"", ctype
        last = f"HTTP {resp.status_code} {ctype}"
        if resp.status_code in (429, 500, 502, 503, 504):
            time.sleep(min(60, 2 ** attempt))
            continue
        return resp.status_code, resp.content or b"", ctype
    raise RuntimeError(last or "请求失败")


def save_atomic(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".part")
    tmp.write_bytes(data)
    tmp.replace(path)


def fetch_pdf(session: requests.Session, limiter: RateLimiter, pmcid: str) -> tuple[str, bytes]:
    status, data, _ctype = request_bytes(
        session, limiter, EUROPEPMC_PDF, {"pmcid": pmcid}
    )
    if status == 200 and is_pdf(data):
        return "ok", data
    if status in (404, 410):
        return "no_oa", b""
    if data and b"error" in data[:200].lower():
        return "no_oa", b""
    return "fail", b""


def fetch_xml(session: requests.Session, limiter: RateLimiter, pmcid: str,
              api_key: str) -> tuple[str, bytes]:
    status, data, _ctype = request_bytes(
        session, limiter, EUROPEPMC_XML.format(pmcid=pmcid)
    )
    if status == 200 and is_xml(data) and b"<error" not in data[:300].lower():
        return "ok", data
    params = {
        "db": "pmc",
        "id": pmcid.replace("PMC", ""),
        "retmode": "xml",
        "rettype": "xml",
        "tool": TOOL,
    }
    if api_key:
        params["api_key"] = api_key
    if EMAIL:
        params["email"] = EMAIL
    status, data, _ctype = request_bytes(session, limiter, EFETCH_URL, params)
    if status == 200 and is_xml(data):
        lowered = data[:500].lower()
        if b"<error" in lowered or b"cannot get document" in lowered:
            return "no_oa", b""
        return "ok", data
    if status in (404, 410):
        return "no_oa", b""
    return "fail", b""


def write_summary_csv(path: Path, store: Store) -> None:
    rows = store.conn.execute(
        "SELECT pmcid, kind, status, path, detail, updated_at FROM downloads ORDER BY pmcid, kind"
    ).fetchall()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["PMCID", "类型", "状态", "路径", "说明", "更新时间"])
        writer.writerows(rows)


def download_one(session: requests.Session, limiter: RateLimiter, store: Store,
                 out_dir: Path, pmcid: str, kinds: list[str], resume: bool,
                 api_key: str, stats: dict[str, int]) -> None:
    for kind in kinds:
        dest = out_dir / kind / f"{pmcid}.{'pdf' if kind == 'pdf' else 'xml'}"
        prev = store.status(pmcid, kind) if resume else None
        if resume and prev == "ok" and dest.exists() and dest.stat().st_size > 0:
            stats["skip"] += 1
            continue
        if resume and dest.exists() and dest.stat().st_size > 0:
            store.mark(pmcid, kind, "ok", str(dest), "已有文件")
            stats["skip"] += 1
            continue
        try:
            if kind == "pdf":
                status, data = fetch_pdf(session, limiter, pmcid)
            else:
                status, data = fetch_xml(session, limiter, pmcid, api_key)
        except Exception as exc:
            store.mark(pmcid, kind, "fail", "", str(exc))
            stats["fail"] += 1
            log(f"失败 {pmcid}/{kind}: {exc}")
            continue
        if status == "ok":
            save_atomic(dest, data)
            store.mark(pmcid, kind, "ok", str(dest), f"{len(data)} bytes")
            stats["ok"] += 1
            log(f"完成 {pmcid}/{kind} -> {dest.name} ({len(data)} bytes)")
        elif status == "no_oa":
            store.mark(pmcid, kind, "no_oa", "", "无开放获取全文")
            stats["no_oa"] += 1
            log(f"无 OA {pmcid}/{kind}")
        else:
            store.mark(pmcid, kind, "fail", "", "下载失败")
            stats["fail"] += 1
            log(f"失败 {pmcid}/{kind}")


def process_files(files: list[Path], out_dir: Path, kinds: list[str], resume: bool,
                  api_key: str, limit: int) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    store = Store(out_dir / "_pmc_download.sqlite")
    session = requests.Session()
    limiter = RateLimiter(MIN_INTERVAL)
    stats = {"ok": 0, "skip": 0, "no_oa": 0, "fail": 0}
    processed = 0
    try:
        for index, path in enumerate(files, start=1):
            log(f"开始处理 ({index}/{len(files)}): {path.name}")
            batch: list[str] = []
            file_new = 0
            for pmcid in iter_pmcids(path):
                if limit > 0 and processed >= limit:
                    break
                # 续跑时若各格式都已完成，直接跳过，不进下载批次
                if resume and all(
                    store.status(pmcid, kind) == "ok"
                    and (out_dir / kind / f"{pmcid}.{'pdf' if kind == 'pdf' else 'xml'}").exists()
                    for kind in kinds
                ):
                    stats["skip"] += len(kinds)
                    processed += 1
                    continue
                batch.append(pmcid)
                file_new += 1
                processed += 1
                if len(batch) >= BATCH_SIZE:
                    for item in batch:
                        download_one(
                            session, limiter, store, out_dir, item, kinds, resume, api_key, stats
                        )
                    log(f"  {path.name} 已处理本文件累计新 PMCID {file_new}")
                    batch.clear()
                if limit > 0 and processed >= limit:
                    break
            for item in batch:
                download_one(
                    session, limiter, store, out_dir, item, kinds, resume, api_key, stats
                )
            log(f"完成表 {path.name}：本文件扫到待处理/续跑相关 PMCID {file_new} 个")
            if limit > 0 and processed >= limit:
                log(f"已达 --limit {limit}，停止后续文件")
                break
        write_summary_csv(out_dir / "_pmc_download_log.csv", store)
        log(
            f"结束：成功 {stats['ok']}，跳过 {stats['skip']}，"
            f"无OA {stats['no_oa']}，失败 {stats['fail']}。日志: {out_dir / '_pmc_download_log.csv'}"
        )
    finally:
        store.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="按题录 Excel 中的 PMCID 下载 PMC 开放获取 PDF/XML")
    parser.add_argument("--excel", default="", help="指定单个题录 Excel 文件（优先于 --excel-dir）")
    parser.add_argument("--excel-dir", default=DEFAULT_EXCEL_DIR, help="题录 Excel 所在文件夹")
    parser.add_argument("--out", default=DEFAULT_OUT_DIR, help="全文输出目录；未指定时：单文件则用其同级 _全文，目录模式用 excel-dir/_全文")
    parser.add_argument(
        "--format",
        default="both",
        choices=("both", "pdf", "xml"),
        help="下载类型：both / pdf / xml",
    )
    parser.add_argument("--resume", action="store_true", help="跳过已成功下载的 PMCID")
    parser.add_argument("--api-key", default=API_KEY, help="NCBI API Key，仅用于 XML 回退 EFetch")
    parser.add_argument("--limit", type=int, default=0, help="最多处理 N 个 PMCID，0 表示全部")
    parser.add_argument("--batch-size", type=int, default=BATCH_SIZE, help="边读边下的批次大小")
    return parser


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    args = build_parser().parse_args()
    global BATCH_SIZE
    if args.batch_size and args.batch_size > 0:
        BATCH_SIZE = args.batch_size
    files = resolve_excel_inputs(args.excel, args.excel_dir)
    if str(args.out).strip():
        out_dir = Path(args.out)
    elif args.excel.strip():
        out_dir = files[0].parent / "_全文"
    else:
        out_dir = Path(args.excel_dir) / "_全文"
    kinds = ["pdf", "xml"] if args.format == "both" else [args.format]
    log(f"输入: {files[0] if len(files) == 1 else f'{len(files)} 个 Excel'}")
    log(f"输出目录: {out_dir}")
    log(f"下载类型: {', '.join(kinds)}；批次 {BATCH_SIZE}；续跑={'是' if args.resume else '否'}")
    process_files(files, out_dir, kinds, args.resume, args.api_key.strip(), args.limit)


if __name__ == "__main__":
    main()
